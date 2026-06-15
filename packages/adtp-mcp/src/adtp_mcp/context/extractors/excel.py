from __future__ import annotations

import time
from pathlib import Path

from adtp_mcp.context import ExtractionResult


class ExcelExtractor:
    @staticmethod
    def extract(file_path: str | Path) -> ExtractionResult:
        start = time.perf_counter()
        path = Path(file_path)
        ext = path.suffix.lower()

        schema: dict = {}
        summary = ""
        agent_hint = ""
        error: str | None = None

        try:
            if ext == ".xls":
                import xlrd

                wb = xlrd.open_workbook(str(path))
                sheet_names = wb.sheet_names()
                headers: dict[str, list[str]] = {}
                row_count = 0
                column_count = 0
                has_formulas = False
                for name in sheet_names:
                    ws = wb.sheet_by_name(name)
                    if ws.nrows > 0:
                        headers[name] = [
                            str(ws.cell_value(0, c)) for c in range(ws.ncols)
                        ]
                    row_count += ws.nrows
                    column_count = max(column_count, ws.ncols)

                schema = {
                    "sheet_names": sheet_names,
                    "row_count": row_count,
                    "column_count": column_count,
                    "headers": headers,
                    "has_formulas": has_formulas,
                    "has_charts": False,
                    "detected_data_types": {},
                }
            else:
                import openpyxl

                wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
                sheet_names = wb.sheetnames
                headers: dict[str, list[str]] = {}
                row_count = 0
                column_count = 0
                detected_types: dict[str, str] = {}

                for name in sheet_names:
                    ws = wb[name]
                    rows_iter = ws.iter_rows(max_row=2, values_only=True)
                    first_row = next(rows_iter, None)
                    if first_row:
                        headers[name] = [
                            str(v) if v is not None else "" for v in first_row
                        ]
                        column_count = max(column_count, len(first_row))
                        second_row = next(rows_iter, None)
                        if second_row:
                            for i, val in enumerate(second_row):
                                col_name = (
                                    headers[name][i]
                                    if i < len(headers[name])
                                    else f"col_{i}"
                                )
                                detected_types[col_name] = (
                                    type(val).__name__ if val is not None else "null"
                                )

                    ws_row_count = sum(
                        1 for _ in ws.iter_rows(max_row=100, values_only=True)
                    )
                    row_count += ws_row_count

                schema = {
                    "sheet_names": sheet_names,
                    "row_count": row_count,
                    "column_count": column_count,
                    "headers": headers,
                    "has_formulas": False,
                    "has_charts": False,
                    "detected_data_types": detected_types,
                }
                wb.close()

            main_sheet = sheet_names[0] if sheet_names else "Unknown"
            summary = f"Excel workbook with {len(sheet_names)} sheet(s), ~{row_count} rows, {column_count} columns"
            agent_hint = f"Primary processing target: sheet '{main_sheet}' — contains {len(headers.get(main_sheet, []))} columns"

        except Exception as e:
            error = str(e)
            summary = "Failed to extract Excel metadata"
            agent_hint = "Binary file — process by filename and size only"

        elapsed = int((time.perf_counter() - start) * 1000)
        return ExtractionResult(
            schema=schema,
            summary=summary[:200],
            agent_hint=agent_hint,
            extraction_ms=elapsed,
            error=error,
        )
