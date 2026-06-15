from __future__ import annotations

import time
from pathlib import Path

from adtp_mcp.context import ExtractionResult


class WordExtractor:
    @staticmethod
    def extract(file_path: str | Path) -> ExtractionResult:
        start = time.perf_counter()
        schema: dict = {}
        summary = ""
        agent_hint = ""
        error: str | None = None

        try:
            from docx import Document

            doc = Document(str(file_path))

            headings: list[str] = []
            for p in doc.paragraphs:
                if p.style.name.startswith("Heading"):
                    headings.append(p.text)

            word_count = sum(
                len(p.text.split()) for p in doc.paragraphs if p.text.strip()
            )
            paragraph_count = len([p for p in doc.paragraphs if p.text.strip()])
            table_count = len(doc.tables)
            has_images = any(
                hasattr(paragraph, "_element")
                and paragraph._element.findall(
                    ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing"
                )
                for paragraph in doc.paragraphs
            )
            core_props = doc.core_properties or None
            author = core_props.author if core_props and core_props.author else ""
            last_modified = str(core_props.modified or "") if core_props else ""

            schema = {
                "word_count": word_count,
                "paragraph_count": paragraph_count,
                "heading_count": len(headings),
                "headings": headings,
                "table_count": table_count,
                "has_images": has_images,
                "author": author or "",
                "last_modified": last_modified,
            }

            summary = f"Word document with ~{word_count} words, {len(headings)} heading(s), {table_count} table(s)"
            agent_hint = f"Document has {len(headings)} section(s) — extract key points from each section"

        except Exception as e:
            error = str(e)
            summary = "Failed to extract Word document metadata"
            agent_hint = "Binary file — process by filename and size only"

        elapsed = int((time.perf_counter() - start) * 1000)
        return ExtractionResult(
            schema=schema,
            summary=summary[:200],
            agent_hint=agent_hint,
            extraction_ms=elapsed,
            error=error,
        )
